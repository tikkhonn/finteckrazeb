#!/usr/bin/env python3
"""
Собирает pay-and-go.pptx из HTML-деки: снимает каждый слайд 1920×1080 и кладёт на страницу PPTX.

Требования (один раз):
  pip install python-pptx playwright
  playwright install chromium

Запуск из корня репозитория или из presentation/:
  python3 presentation/export-pptx.py
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
HTML = HERE / "pay-and-go.html"
OUT = HERE / "pay-and-go.pptx"
SLIDE_COUNT = 14
VIEWPORT = {"width": 1920, "height": 1080}
WAIT_MS = 2200  # анимации reveal


def main() -> int:
    if not HTML.is_file():
        print(f"Не найден {HTML}", file=sys.stderr)
        return 1

    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("Установите: pip install playwright && playwright install chromium", file=sys.stderr)
        return 1

    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        print("Установите: pip install python-pptx", file=sys.stderr)
        return 1

    shots: list[Path] = []
    base_url = HTML.resolve().as_uri()

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport=VIEWPORT)
        with tempfile.TemporaryDirectory(prefix="paygo-slides-") as tmp:
            tmp_path = Path(tmp)
            page.goto(f"{base_url}#1", wait_until="networkidle")
            page.wait_for_function("() => window.__payGoDeck")

            for i in range(1, SLIDE_COUNT + 1):
                # Смена только #N не перезагружает file:// — переключаем слайд явно
                page.evaluate("(idx) => window.__payGoDeck.show(idx)", i - 1)
                page.wait_for_timeout(WAIT_MS)
                path = tmp_path / f"slide-{i:02d}.png"
                page.locator(".deck-stage").screenshot(path=str(path))
                shots.append(path)
            browser.close()

            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            blank = prs.slide_layouts[6]

            for path in shots:
                slide = prs.slides.add_slide(blank)
                slide.shapes.add_picture(
                    str(path), 0, 0, width=prs.slide_width, height=prs.slide_height
                )

            prs.save(OUT)

    print(f"Готово: {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
